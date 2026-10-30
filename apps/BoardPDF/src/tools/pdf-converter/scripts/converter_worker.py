import json
import os
import re
import sys
import tempfile

def convert_to_docx(input_pdf, output_file):
    cv = None
    try:
        from pdf2docx import Converter
        cv = Converter(input_pdf)
        cv.convert(output_file, start=0, end=None)
        cv.close()
        cv = None
        if os.path.exists(output_file) and os.path.getsize(output_file) > 0:
            return
    except Exception:
        pass
    finally:
        if cv:
            try:
                cv.close()
            except Exception:
                pass

    import docx
    from docx.shared import Inches, Pt
    import pymupdf

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    pdf = pymupdf.open(input_pdf)
    for p_idx, page in enumerate(pdf):
        tables = page.find_tables()
        table_rects = []
        if tables and tables.tables:
            for tab in tables.tables:
                table_rects.append(tab.bbox)

        blocks = page.get_text("blocks")
        for b in blocks:
            text = b[4].strip()
            if not text:
                continue
            is_in_table = any(
                b[0] >= tr[0] and b[1] >= tr[1] and b[2] <= tr[2] and b[3] <= tr[3]
                for tr in table_rects
            )
            if is_in_table:
                continue
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(11)

        if tables and tables.tables:
            for tab in tables.tables:
                data = tab.extract()
                if not data:
                    continue
                docx_table = doc.add_table(rows=len(data), cols=len(data[0]))
                docx_table.style = 'Table Grid'
                for r_idx, row in enumerate(data):
                    for c_idx, cell_value in enumerate(row):
                        docx_table.cell(r_idx, c_idx).text = str(cell_value or '').strip()

        if p_idx < len(pdf) - 1:
            doc.add_page_break()

    pdf.close()
    doc.save(output_file)

def convert_to_pptx(input_pdf, output_file):
    from pptx import Presentation
    from pptx.util import Inches
    import pymupdf

    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]

    doc = pymupdf.open(input_pdf)
    for page in doc:
        rect = page.rect
        w_in = max(rect.width / 72.0, 1.0)
        h_in = max(rect.height / 72.0, 1.0)

        prs.slide_width = Inches(w_in)
        prs.slide_height = Inches(h_in)
        slide = prs.slides.add_slide(blank_slide_layout)

        pix = page.get_pixmap(dpi=150)
        fd, tmp_path = tempfile.mkstemp(suffix='.png')
        os.close(fd)

        try:
            pix.save(tmp_path)
            slide.shapes.add_picture(tmp_path, Inches(0), Inches(0), width=Inches(w_in), height=Inches(h_in))
        finally:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass

    doc.close()
    prs.save(output_file)

def convert_to_xlsx(input_pdf, output_file, options=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    import pdfplumber

    options = options or {}
    single_sheet = options.get('excelSheets', 'single') == 'single'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja 1" if single_sheet else "Página 1"

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    current_row = 1
    with pdfplumber.open(input_pdf) as pdf:
        for p_idx, page in enumerate(pdf.pages):
            if single_sheet:
                current_ws = ws
            else:
                current_ws = ws if p_idx == 0 else wb.create_sheet(title=f"Página {p_idx + 1}")
                current_row = 1

            tables = page.extract_tables({
                "vertical_strategy": "lines",
                "horizontal_strategy": "lines",
                "snap_tolerance": 3,
                "join_tolerance": 3,
            })
            if not tables:
                tables = page.extract_tables({
                    "vertical_strategy": "text",
                    "horizontal_strategy": "text",
                    "snap_tolerance": 4,
                })

            if tables:
                for table in tables:
                    for r_i, row in enumerate(table):
                        for c_i, cell in enumerate(row, start=1):
                            val = (cell or "").strip()
                            c = current_ws.cell(row=current_row, column=c_i, value=val)
                            c.font = header_font if r_i == 0 else regular_font
                            c.border = thin_border
                            if r_i == 0:
                                c.fill = header_fill
                        current_row += 1
                    current_row += 1
            else:
                text = page.extract_text(layout=True)
                if text:
                    for line in text.splitlines():
                        trimmed = line.strip()
                        if not trimmed:
                            continue
                        parts = [p.strip() for p in re.split(r'\s{2,}', trimmed) if p.strip()]
                        if not parts:
                            parts = [trimmed]
                        for c_i, cell_val in enumerate(parts, start=1):
                            c = current_ws.cell(row=current_row, column=c_i, value=cell_val)
                            c.font = regular_font
                            c.border = thin_border
                        current_row += 1

            for col in current_ws.columns:
                max_len = 0
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                current_ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    wb.save(output_file)

def convert_to_jpg(input_pdf, output_pattern, options=None):
    import pymupdf
    options = options or {}
    pdf = pymupdf.open(input_pdf)
    mode = options.get('jpgMode', 'pages')
    quality = options.get('jpgQuality', 'normal')
    dpi = 300 if quality == 'high' else 150

    if mode == 'extract':
        img_idx = 1
        for page in pdf:
            for img in page.get_images(full=True):
                xref = img[0]
                try:
                    base_image = pdf.extract_image(xref)
                    image_bytes = base_image["image"]
                    image_ext = base_image.get("ext", "jpg") or "jpg"
                    if image_ext.lower() not in ['jpg', 'jpeg']:
                        pix = pymupdf.Pixmap(image_bytes)
                        if pix.n >= 5:
                            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                        out_name = f"{output_pattern}-img_{img_idx:03d}.jpg"
                        pix.save(out_name)
                    else:
                        out_name = f"{output_pattern}-img_{img_idx:03d}.jpg"
                        with open(out_name, "wb") as f_out:
                            f_out.write(image_bytes)
                    img_idx += 1
                except Exception:
                    pass
        if img_idx == 1:
            total_pages = len(pdf)
            for p_idx, page in enumerate(pdf, start=1):
                pix = page.get_pixmap(dpi=dpi)
                out_name = f"{output_pattern}-{p_idx:02d}.jpg" if total_pages > 1 else f"{output_pattern}.jpg"
                pix.save(out_name)
    else:
        total_pages = len(pdf)
        for p_idx, page in enumerate(pdf, start=1):
            pix = page.get_pixmap(dpi=dpi)
            out_name = f"{output_pattern}-{p_idx:02d}.jpg" if total_pages > 1 else f"{output_pattern}.jpg"
            pix.save(out_name)
    pdf.close()

def main():
    if len(sys.argv) < 4:
        print(json.dumps({'error': 'Uso: python converter_worker.py <format> <input> <output>', 'success': False}))
        sys.exit(1)

    target_format = sys.argv[1].lower()
    input_pdf = sys.argv[2]
    output_file = sys.argv[3]
    options = {}
    if len(sys.argv) > 4:
        try:
            options = json.loads(sys.argv[4])
        except Exception:
            pass

    if not os.path.exists(input_pdf):
        print(json.dumps({'error': 'Archivo de entrada no encontrado', 'success': False}))
        sys.exit(1)

    try:
        if target_format == 'docx':
            convert_to_docx(input_pdf, output_file)
        elif target_format == 'pptx':
            convert_to_pptx(input_pdf, output_file)
        elif target_format == 'xlsx':
            convert_to_xlsx(input_pdf, output_file, options)
        elif target_format == 'jpg' or target_format == 'jpeg':
            convert_to_jpg(input_pdf, output_file, options)
        else:
            print(json.dumps({'error': f'Formato no soportado: {target_format}', 'success': False}))
            sys.exit(1)

        print(json.dumps({'output': output_file, 'success': True}))
    except Exception as e:
        print(json.dumps({'error': str(e), 'success': False}))
        sys.exit(1)

if __name__ == '__main__':
    main()
